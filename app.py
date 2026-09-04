import json
import os
import tempfile
from typing import List

import streamlit as st
from google import genai
from google.genai import types
from pydantic import BaseModel, Field
from pypdf import PdfReader
from docx import Document


# -----------------------------
# App configuration
# -----------------------------
st.set_page_config(
    page_title="AI Resume ATS Analyzer",
    page_icon="📄",
    layout="wide",
)

MODEL_NAME = "gemini-2.5-flash"


# -----------------------------
# Structured Gemini response
# -----------------------------
class ScoreItem(BaseModel):
    category: str
    score: int = Field(ge=0, le=100)
    weight: int = Field(ge=0, le=100)
    explanation: str


class ResumeAnalysis(BaseModel):
    overall_score: int = Field(ge=0, le=100)
    verdict: str
    summary: str
    score_breakdown: List[ScoreItem]
    strengths: List[str]
    improvements: List[str]
    missing_keywords: List[str]
    detected_keywords: List[str]
    formatting_issues: List[str]
    action_plan: List[str]


# -----------------------------
# Helpers
# -----------------------------
def get_api_key() -> str:
    """Read the Gemini key from Streamlit secrets or environment variables."""
    try:
        key = st.secrets.get("GEMINI_API_KEY")
    except Exception:
        key = None

    key = key or os.getenv("GEMINI_API_KEY") or os.getenv("GOOGLE_API_KEY")

    if not key:
        raise RuntimeError(
            "Gemini API key not found. Add GEMINI_API_KEY to "
            ".streamlit/secrets.toml locally or Streamlit Cloud Secrets."
        )

    return key


def extract_pdf_text(uploaded_file) -> str:
    reader = PdfReader(uploaded_file)
    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""
        pages.append(text)

    return "\n".join(pages).strip()


def extract_docx_text(uploaded_file) -> str:
    document = Document(uploaded_file)
    parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def extract_text(uploaded_file) -> str:
    file_type = uploaded_file.type
    name = uploaded_file.name.lower()

    if file_type == "application/pdf" or name.endswith(".pdf"):
        return extract_pdf_text(uploaded_file)

    if (
        file_type
        == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        or name.endswith(".docx")
    ):
        return extract_docx_text(uploaded_file)

    if file_type.startswith("text/") or name.endswith(".txt"):
        return uploaded_file.getvalue().decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload PDF, DOCX, or TXT.")


def build_prompt(resume_text: str, job_description: str) -> str:
    job_context = (
        f"""
TARGET JOB DESCRIPTION:
{job_description}
"""
        if job_description.strip()
        else """
TARGET JOB DESCRIPTION:
Not provided. Evaluate general ATS readiness and professional resume quality.
Do NOT pretend to have exact job-specific keyword matching without a job description.
"""
    )

    return f"""
You are an expert ATS resume evaluator and career coach.

Analyze the candidate resume below. Produce a practical, honest ATS-readiness
assessment from 0 to 100.

Important:
- This is an estimated ATS-readiness score, not a score from a specific ATS vendor.
- Do not invent experience, education, certifications, skills, employers, dates,
  achievements, or keywords that are not present.
- If a target job description is supplied, compare the resume against it.
- If no job description is supplied, evaluate general ATS best practices.
- Consider keyword relevance, section structure, measurable achievements,
  skills, clarity, contact information, and machine-readable formatting.
- Give concrete improvements, not vague advice.
- Keep the score breakdown internally consistent with the overall score.
- Missing keywords should only be reported when they are genuinely relevant to
  the supplied job description or, when no job description exists, clearly useful
  for the resume's apparent target role.
- Never recommend keyword stuffing.

{job_context}

RESUME:
{resume_text}
"""


def analyze_resume(resume_text: str, job_description: str) -> ResumeAnalysis:
    client = genai.Client(api_key=get_api_key())

    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
            response_schema=ResumeAnalysis,
        ),
    )

    if not response.text:
        raise RuntimeError("Gemini returned an empty response.")

    return ResumeAnalysis.model_validate_json(response.text)


def score_label(score: int) -> str:
    if score >= 85:
        return "Excellent ATS readiness"
    if score >= 70:
        return "Good ATS readiness"
    if score >= 55:
        return "Needs improvement"
    return "High improvement needed"


# -----------------------------
# UI
# -----------------------------
st.title("📄 AI Resume ATS Analyzer")
st.markdown(
    "Upload your resume and get an **estimated ATS score**, strengths, "
    "missing keywords, formatting issues, and an actionable improvement plan."
)

st.info(
    "Tip: Add the job description for a more useful, job-specific keyword and ATS analysis."
)

col1, col2 = st.columns([1, 1])

with col1:
    uploaded_file = st.file_uploader(
        "Upload your resume",
        type=["pdf", "docx", "txt"],
        help="Supported formats: PDF, DOCX, TXT.",
    )

with col2:
    job_description = st.text_area(
        "Target job description (optional)",
        height=170,
        placeholder="Paste the job description here for job-specific ATS matching...",
    )

analyze_button = st.button(
    "🔍 Analyze Resume",
    type="primary",
    use_container_width=True,
)

if analyze_button:
    if uploaded_file is None:
        st.warning("Please upload a resume first.")
        st.stop()

    with st.spinner("Reading your resume and analyzing ATS readiness..."):
        try:
            resume_text = extract_text(uploaded_file)

            if len(resume_text.strip()) < 80:
                st.error(
                    "Very little readable text was found. If this is a scanned/image-only PDF, "
                    "please upload a text-based PDF or DOCX version."
                )
                st.stop()

            # Keep extremely large extracted text from creating an unnecessarily
            # large prompt while retaining the beginning and end of the resume.
            max_chars = 60000
            if len(resume_text) > max_chars:
                resume_text = (
                    resume_text[:45000]
                    + "\n\n[Middle of resume truncated for processing]\n\n"
                    + resume_text[-15000:]
                )

            result = analyze_resume(resume_text, job_description)

        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.stop()

    st.success("Analysis complete!")

    # Main score
    score_col1, score_col2 = st.columns([1, 2])

    with score_col1:
        st.metric("Estimated ATS Score", f"{result.overall_score}/100")
        st.progress(result.overall_score / 100)
        st.caption(score_label(result.overall_score))

    with score_col2:
        st.subheader(result.verdict)
        st.write(result.summary)

    st.divider()

    # Score breakdown
    st.subheader("📊 Score Breakdown")

    for item in result.score_breakdown:
        st.markdown(f"**{item.category} — {item.score}/100** (weight: {item.weight}%)")
        st.progress(item.score / 100)
        st.caption(item.explanation)

    st.divider()

    left, right = st.columns(2)

    with left:
        st.subheader("✅ Strengths")
        if result.strengths:
            for item in result.strengths:
                st.markdown(f"- {item}")
        else:
            st.write("No major strengths were detected.")

        st.subheader("🔑 Detected Keywords")
        if result.detected_keywords:
            st.write(", ".join(result.detected_keywords))
        else:
            st.write("No strong keywords were detected.")

    with right:
        st.subheader("⚠️ Improvements")
        if result.improvements:
            for item in result.improvements:
                st.markdown(f"- {item}")
        else:
            st.write("No major improvements were detected.")

        st.subheader("❌ Missing / Useful Keywords")
        if result.missing_keywords:
            for item in result.missing_keywords:
                st.markdown(f"- {item}")
        else:
            st.write("No important missing keywords were identified.")

    st.divider()

    st.subheader("🧾 Formatting & ATS Issues")
    if result.formatting_issues:
        for item in result.formatting_issues:
            st.markdown(f"- {item}")
    else:
        st.write("No major formatting issues were detected.")

    st.subheader("🚀 Action Plan")
    for index, item in enumerate(result.action_plan, start=1):
        st.markdown(f"**{index}.** {item}")

    st.caption(
        "This tool provides an AI-generated ATS-readiness estimate. "
        "Different applicant tracking systems use different parsing and ranking rules, "
        "so the score should be treated as guidance rather than a guaranteed hiring outcome."
    )

st.divider()
st.caption("Built with Streamlit + Gemini 2.5 Flash")
